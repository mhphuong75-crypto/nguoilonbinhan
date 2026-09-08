#!/usr/bin/env python3
"""nguoilonbinhan.meschool - boc checklist SUPA trong kho 'kho-tho'.

Moi tep SUPA:
  sheet 'Report'          : STT | Ma Questionnaire | Ten Questionnaire
  sheet '<Ma>' (moi bo)   : dong tieu de co 'Nguoi tra loi'; cot 0..7 la thong tin,
                            cot 8+ la tung cau hoi, dap an 'Co' / 'Khong' / ...

Ghi ra:
  supa_phieu  : 1 dong = 1 lan lam checklist (ty le dat)
  supa_truot  : 1 dong = 1 cau KHONG dat  -> de biet cau nao hay truot nhat
  nguoi_phu_trach : tu dong cap nhat ten that cua QLVH tung campus
"""
import io, os, re, json, datetime, urllib.request

import openpyxl

SB_URL = os.environ["SUPABASE_URL"].rstrip("/")
SB_KEY = os.environ["SUPABASE_KEY"]
BUCKET = "kho-tho"
PREFIX = "supa/"
DAT = {"có", "co", "đạt", "dat", "yes", "ok"}


def goi(url, method="GET", body=None, headers=None, raw=False):
    h = {"apikey": SB_KEY, "Authorization": "Bearer " + SB_KEY}
    h.update(headers or {})
    data = None
    if body is not None:
        data = json.dumps(body).encode()
        h["Content-Type"] = "application/json"
    req = urllib.request.Request(url, data=data, headers=h, method=method)
    with urllib.request.urlopen(req, timeout=240) as r:
        return r.read() if raw else r.read().decode()


def ghi(bang, dong, xung_dot):
    for i in range(0, len(dong), 1000):
        goi(f"{SB_URL}/rest/v1/{bang}?on_conflict={xung_dot}", "POST", dong[i:i + 1000],
            {"Prefer": "resolution=merge-duplicates,return=minimal"})


def danh_sach_tep():
    out = goi(f"{SB_URL}/storage/v1/object/list/{BUCKET}", "POST",
              {"prefix": PREFIX, "limit": 200,
               "sortBy": {"column": "created_at", "order": "asc"}})
    return [PREFIX + o["name"] for o in json.loads(out) if o.get("name", "").endswith(".xlsx")]


def chu(v, n=400):
    if v is None:
        return None
    s = str(v).strip()
    return s[:n] or None


def ngay_tu(v):
    if v is None or v == "":
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return str(v.date() if isinstance(v, datetime.datetime) else v)
    s = str(v).strip()
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return m.group(0)
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        try:
            return str(datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1))))
        except ValueError:
            return None
    return None


def tach_ten(v):
    """'Phan Kieu Han (QLVH Nguyen Trong Tuyen)' -> ('Phan Kieu Han', 'QLVH', 'Nguyen Trong Tuyen')"""
    s = chu(v, 200) or ""
    m = re.match(r"^(.*?)\s*\((.*)\)\s*$", s)
    if not m:
        return s or None, None, None
    ten, trong = m.group(1).strip(), m.group(2).strip()
    m2 = re.match(r"^([A-Za-zÀ-ỹ\.]+)\s+(.*)$", trong)
    if m2:
        return ten or None, m2.group(1), m2.group(2)
    return ten or None, trong or None, None


def don_dia_diem(v):
    s = chu(v, 120) or ""
    return re.sub(r"^[\s\-–—]+", "", s).strip() or None


def boc_bo(ws, bo_ma, bo_ten, tep):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    i_td = None
    for i, r in enumerate(rows[:20]):
        v = [chu(c, 60) or "" for c in r]
        if any("người trả lời" in x.lower() for x in v):
            i_td, td = i, v
            break
    if i_td is None:
        return [], []
    m = {t.strip().lower(): j for j, t in enumerate(td) if t and t.strip()}
    c_id = m.get("id")
    c_ng = m.get("người trả lời")
    c_dd = m.get("địa điểm")
    c_ngay = m.get("ngày thực hiện")
    c_ql = m.get("quản lý trực tiếp")
    dau_cau = max([j for j in (c_id, c_ng, c_dd, c_ngay, c_ql, m.get("mã địa điểm"),
                               m.get("mã nhân viên"), m.get("stt")) if j is not None]) + 1
    cau_ten = {j: (td[j] or "")[:200] for j in range(dau_cau, len(td)) if td[j]}

    phieu, truot, thay = [], [], set()
    for r in rows[i_td + 1:]:
        pid = chu(r[c_id], 60) if c_id is not None and c_id < len(r) else None
        if not pid or pid in thay:
            continue
        thay.add(pid)
        ten, vai, cs = tach_ten(r[c_ng] if c_ng is not None and c_ng < len(r) else None)
        dd = don_dia_diem(r[c_dd] if c_dd is not None and c_dd < len(r) else None) or cs
        ngay = ngay_tu(r[c_ngay] if c_ngay is not None and c_ngay < len(r) else None)
        n_cau = n_dat = 0
        for j, ten_cau in cau_ten.items():
            if j >= len(r):
                continue
            da = chu(r[j], 200)
            if not da:
                continue
            n_cau += 1
            if da.strip().lower() in DAT:
                n_dat += 1
            else:
                truot.append({"khoa": f"{pid}_{j}", "phieu_id": pid, "bo_ma": bo_ma,
                              "ngay": ngay, "dia_diem": dd, "nguoi_tra_loi": ten,
                              "cau": ten_cau, "dap_an": da[:200]})
        if not n_cau:
            continue
        phieu.append({"id": pid, "bo_ma": bo_ma, "bo_ten": bo_ten, "ngay": ngay,
                      "dia_diem": dd, "nguoi_tra_loi": ten, "vai_tro": vai,
                      "quan_ly": chu(r[c_ql], 120) if c_ql is not None and c_ql < len(r) else None,
                      "so_cau": n_cau, "so_dat": n_dat,
                      "ty_le": round(100.0 * n_dat / n_cau, 1), "tep": tep})
    return phieu, truot


def cap_nhat_nguoi(phieu):
    """Lay ten QLVH moi nhat cho tung campus."""
    tot = {}
    for p in phieu:
        dd, ten, vai, ngay = p.get("dia_diem"), p.get("nguoi_tra_loi"), p.get("vai_tro"), p.get("ngay") or ""
        if not dd or not ten or not vai or "QLVH" not in vai.upper():
            continue
        if dd not in tot or ngay > tot[dd]["ngay"]:
            tot[dd] = {"ngay": ngay, "ten": ten, "vai": vai}
    dong = [{"dia_diem": k, "ten_that": v["ten"], "vai_tro": v["vai"], "ngay_thay": v["ngay"] or None}
            for k, v in tot.items()]
    if dong:
        ghi("nguoi_phu_trach_supa", dong, "dia_diem")
    return len(dong)


def main():
    teps = danh_sach_tep()
    print(f"Tim thay {len(teps)} tep SUPA")
    tong_p = tong_t = 0
    tat_ca_phieu = []
    for ten in teps:
        try:
            b = goi(f"{SB_URL}/storage/v1/object/{BUCKET}/{ten}", raw=True)
            wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
            ten_bo = {}
            if "Report" in wb.sheetnames:
                for r in wb["Report"].iter_rows(values_only=True):
                    if r and len(r) >= 3 and chu(r[1]) and str(r[1]).strip() != "Mã Questionnaire":
                        ten_bo[str(r[1]).strip()] = chu(r[2], 200)
            for s in wb.sheetnames:
                if s == "Report":
                    continue
                p, t = boc_bo(wb[s], s, ten_bo.get(s), ten)
                if p:
                    ghi("supa_phieu", p, "id")
                    tat_ca_phieu += p
                    tong_p += len(p)
                if t:
                    ghi("supa_truot", t, "khoa")
                    tong_t += len(t)
                print(f"  {ten} :: {s} -> {len(p)} phieu, {len(t)} cau truot")
            wb.close()
        except Exception as e:
            print(f"  {ten}: LOI {type(e).__name__} {e}")
    print(f"Xong. {tong_p} phieu, {tong_t} cau truot.")
    try:
        n = cap_nhat_nguoi(tat_ca_phieu)
        print(f"Cap nhat {n} nguoi phu trach tu SUPA")
    except Exception as e:
        print(f"Nguoi phu trach: LOI {type(e).__name__} {e}")


if __name__ == "__main__":
    main()
