#!/usr/bin/env python3
"""Soi cau truc file trong kho. Chay bang tay khi mot nguon doi dinh dang."""
import io, os, json, urllib.request

import openpyxl

SB = os.environ["SUPABASE_URL"].rstrip("/")
KEY = os.environ["SUPABASE_KEY"]
BUCKET = "kho-tho"
CHI = os.environ.get("CHI_PREFIX", "supa/")
SO_DONG = int(os.environ.get("SO_DONG", "16"))


def goi(url, method="GET", body=None, raw=False):
    h = {"apikey": KEY, "Authorization": "Bearer " + KEY}
    data = None
    if body is not None:
        data = json.dumps(body).encode()
        h["Content-Type"] = "application/json"
    req = urllib.request.Request(url, data=data, headers=h, method=method)
    with urllib.request.urlopen(req, timeout=180) as r:
        return r.read() if raw else r.read().decode()


def moi_nhat(prefix, n=1):
    out = goi(f"{SB}/storage/v1/object/list/{BUCKET}", "POST",
              {"prefix": prefix, "limit": 100,
               "sortBy": {"column": "created_at", "order": "desc"}})
    ds = [prefix + o["name"] for o in json.loads(out) if o.get("name", "").endswith(".xlsx")]
    return ds[:n]


def soi(ten):
    print(f"\n{'='*70}\nTEP: {ten}")
    b = goi(f"{SB}/storage/v1/object/{BUCKET}/{ten}", raw=True)
    print(f"  kich thuoc: {len(b)//1024} KB")
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    print(f"  SO SHEET: {len(wb.sheetnames)} -> {wb.sheetnames}")
    for s in wb.sheetnames:
        ws = wb[s]
        print(f"\n  ### SHEET {s}  ({ws.max_row} dong x {ws.max_column} cot)")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= SO_DONG:
                break
            v = [("" if c is None else str(c))[:30] for c in row[:20]]
            if any(x.strip() for x in v):
                print(f"    r{i}: " + " ~ ".join(v))
    wb.close()


for t in moi_nhat(CHI, 2):
    try:
        soi(t)
    except Exception as e:
        print(f"{t}: LOI {type(e).__name__} {e}")
